import openpyxl
from datetime import date, datetime
from app.database import Base, engine, SessionLocal
from app.models.user import User
from app.models.config_setting import ConfigSetting
from app.models.session import MonitoringSession
from app.models.scan_data import ScanPoint
from app.models.sims import SimsDataset, SimsRecord
from app.models.manual_station import ManualStation
from app.models.identification import IdentificationResult
from app.core.security import get_password_hash
from app.core.identification_engine import IdentificationEngine

def seed():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    
    # Check if admin exists
    if not db.query(User).filter(User.username == "admin").first():
        db.add(User(
            username="admin",
            full_name="Administrator Balmon DIY",
            nip="198501012010011001",
            role="admin",
            password_hash=get_password_hash("admin123")
        ))
    if not db.query(User).filter(User.username == "petugas").first():
        db.add(User(
            username="petugas",
            full_name="Agus Dong (PPFR Ahli Muda)",
            nip="199002022015021002",
            role="petugas",
            password_hash=get_password_hash("petugas123")
        ))
        
    if not db.query(ConfigSetting).filter(ConfigSetting.setting_key == "default").first():
        db.add(ConfigSetting(
            setting_key="default",
            lat_min=-8.213584,
            lat_max=-7.503037,
            lon_min=109.375246,
            lon_max=111.321684,
            default_threshold_dbuvm=50.0,
            default_radius_km=50.0
        ))
    db.commit()
    
    # Read sample data from Excel workbook if present
    xlsm_path = r'g:\My Drive\Google AI Studio\AI Monitoring Balmon Yogyakarta V.1.0\Tools_Auto_Identification_Monitoring_Balmon_Yogyakarta.xlsm'
    try:
        wb = openpyxl.load_workbook(xlsm_path, data_only=True)
        
        # 1. TambahanDataSIMS
        if "TambahanDataSIMS" in wb.sheetnames:
            ws_tam = wb["TambahanDataSIMS"]
            for row in ws_tam.iter_rows(min_row=2, max_row=50, values_only=True):
                if row and row[0] is not None and row[1] is not None:
                    freq = float(str(row[0]).replace(",", "."))
                    name = str(row[1]).strip()
                    lat = float(str(row[2]).replace(",", ".")) if row[2] is not None else -7.733042
                    lon = float(str(row[3]).replace(",", ".")) if row[3] is not None else 110.471837
                    serv = str(row[7]).strip() if len(row) > 7 and row[7] else None
                    subserv = str(row[8]).strip() if len(row) > 8 and row[8] else None
                    emis = str(row[9]).strip() if len(row) > 9 and row[9] else None
                    
                    if not db.query(ManualStation).filter(ManualStation.freq_mhz == freq, ManualStation.client_name == name).first():
                        db.add(ManualStation(
                            freq_mhz=freq, client_name=name, latitude=lat, longitude=lon,
                            service=serv, subservice=subserv, emission_class=emis, status_legality="Legal"
                        ))
            db.commit()
            
        # 2. Sample Dataset SIMS
        sims_ds = db.query(SimsDataset).filter(SimsDataset.dataset_name == "SIMS_BALMON_DIY_MASTER").first()
        if not sims_ds:
            sims_ds = SimsDataset(
                dataset_name="SIMS_BALMON_DIY_MASTER",
                query_date=date(2026, 9, 1),
                uploaded_by="Agus Dong",
                total_records=8,
                is_active=True
            )
            db.add(sims_ds)
            db.commit()
            db.refresh(sims_ds)
            
            sample_sims = [
                SimsRecord(dataset_id=sims_ds.id, freq_mhz=88.2, client_name="PT RADIO RETJO BUNTUNG", latitude=-7.7956, longitude=110.3695, service="PENYIARAN", subservice="RADIO FM", emission_class="F3E", station_name="RETJO BUNTUNG FM", city="KOTA YOGYAKARTA", expiry_date=date(2027, 12, 31)),
                SimsRecord(dataset_id=sims_ds.id, freq_mhz=89.0, client_name="PT RADIO GERONIMO", latitude=-7.7828, longitude=110.3742, service="PENYIARAN", subservice="RADIO FM", emission_class="F3E", station_name="GERONIMO FM", city="KOTA YOGYAKARTA", expiry_date=date(2027, 10, 15)),
                SimsRecord(dataset_id=sims_ds.id, freq_mhz=91.4, client_name="PT RADIO UNISI", latitude=-7.7712, longitude=110.3790, service="PENYIARAN", subservice="RADIO FM", emission_class="F3E", station_name="UNISI FM", city="SLEMAN", expiry_date=date(2026, 11, 20)),
                SimsRecord(dataset_id=sims_ds.id, freq_mhz=93.0, client_name="PT RADIO SWADANA FM", latitude=-7.8010, longitude=110.3550, service="PENYIARAN", subservice="RADIO FM", emission_class="F3E", station_name="SWADANA FM", city="BANTUL", expiry_date=date(2024, 5, 1)), # Expired
                SimsRecord(dataset_id=sims_ds.id, freq_mhz=97.0, client_name="PT RADIO MQ FM", latitude=-7.7560, longitude=110.4080, service="PENYIARAN", subservice="RADIO FM", emission_class="F3E", station_name="MQ FM", city="SLEMAN", expiry_date=date(2027, 8, 30)),
                SimsRecord(dataset_id=sims_ds.id, freq_mhz=100.2, client_name="PT RADIO JOGJA STREAM", latitude=-7.7600, longitude=110.4200, service="PENYIARAN", subservice="RADIO FM", emission_class="F3E", station_name="JOGJA STREAM FM", city="SLEMAN", expiry_date=date(2028, 1, 1)),
                SimsRecord(dataset_id=sims_ds.id, freq_mhz=104.5, client_name="PT RADIO KRISNA FM", latitude=-7.8200, longitude=110.3900, service="PENYIARAN", subservice="RADIO FM", emission_class="F3E", station_name="KRISNA FM", city="BANTUL", expiry_date=date(2023, 12, 31)), # Expired
                SimsRecord(dataset_id=sims_ds.id, freq_mhz=106.9, client_name="PT RADIO ISLAM FM", latitude=-7.7400, longitude=110.4300, service="PENYIARAN", subservice="RADIO FM", emission_class="F3E", station_name="ISLAM FM", city="SLEMAN", expiry_date=date(2027, 4, 15))
            ]
            db.bulk_save_objects(sample_sims)
            db.commit()

        # 3. Create Sample Monitoring Session
        if db.query(MonitoringSession).count() == 0:
            sess = MonitoringSession(
                session_name="Monitoring Rutin Pita FM Yogyakarta 2026",
                upt_name="BALMON SFR KELAS I YOGYAKARTA",
                monitoring_station="STASIUN MONITORING SMSN",
                spt_number="SPT/042/BALMON.34/09/2026",
                officer_name="Agus Dong (PPFR Ahli Muda)",
                monitoring_date=date(2026, 9, 7),
                monitoring_time=datetime.now().time(),
                address="Glondong, Tirtomartani",
                district="Kalasan",
                city="SLEMAN",
                province="D.I. Yogyakarta",
                latitude=-7.733139,
                longitude=110.471667,
                device_type="TCI",
                raw_filename="TCI_FM_Scan_20260907.csv",
                start_frequency_mhz=87.5,
                stop_frequency_mhz=108.0,
                threshold_dbuvm=50.0,
                detection_radius_km=50.0,
                marker_filter="all"
            )
            db.add(sess)
            db.commit()
            db.refresh(sess)
            
            # Generate synthetic FM sweep points with clear peaks
            import numpy as np
            points = []
            cur_f = 87.5
            while cur_f <= 108.0:
                # Base noise floor around 28-35 dBuV/m
                noise = 30.0 + (cur_f % 1.5) * 2.0
                
                # Add signal peaks for known FM stations
                if abs(cur_f - 88.2) < 0.05: noise = 72.4
                elif abs(cur_f - 89.0) < 0.05: noise = 81.2
                elif abs(cur_f - 91.4) < 0.05: noise = 68.9
                elif abs(cur_f - 93.0) < 0.05: noise = 64.1
                elif abs(cur_f - 95.5) < 0.05: noise = 58.7 # Unknown peak
                elif abs(cur_f - 97.0) < 0.05: noise = 75.3
                elif abs(cur_f - 100.2) < 0.05: noise = 84.5
                elif abs(cur_f - 102.8) < 0.05: noise = 62.0 # Unknown peak
                elif abs(cur_f - 104.5) < 0.05: noise = 59.8 # Expired
                elif abs(cur_f - 106.9) < 0.05: noise = 70.2
                
                points.append(ScanPoint(
                    session_id=sess.id,
                    frequency_mhz=round(cur_f, 2),
                    level_dbuvm=round(noise, 1),
                    is_peak=False
                ))
                cur_f += 0.1
                
            db.bulk_save_objects(points)
            db.commit()
            
            # Run identification
            engine_inst = IdentificationEngine(
                monitoring_lat=sess.latitude,
                monitoring_lon=sess.longitude,
                monitoring_date=sess.monitoring_date,
                detection_radius_km=sess.detection_radius_km,
                threshold_dbuvm=sess.threshold_dbuvm
            )
            
            all_pts = db.query(ScanPoint).filter(ScanPoint.session_id == sess.id).order_by(ScanPoint.frequency_mhz.asc()).all()
            pt_dicts = [{"id": p.id, "frequency_mhz": p.frequency_mhz, "level_dbuvm": p.level_dbuvm} for p in all_pts]
            
            all_sims = db.query(SimsRecord).filter(SimsRecord.dataset_id == sims_ds.id).all()
            sims_dicts = [{"id": s.id, "freq_mhz": s.freq_mhz, "client_name": s.client_name, "latitude": s.latitude, "longitude": s.longitude, "service": s.service, "subservice": s.subservice, "emission_class": s.emission_class, "city": s.city, "expiry_date": s.expiry_date} for s in all_sims]
            
            all_man = db.query(ManualStation).filter(ManualStation.is_active == True).all()
            man_dicts = [{"id": m.id, "freq_mhz": m.freq_mhz, "client_name": m.client_name, "latitude": m.latitude, "longitude": m.longitude, "service": m.service, "subservice": m.subservice, "emission_class": m.emission_class, "status_legality": m.status_legality} for m in all_man]
            
            proc = engine_inst.process_session(pt_dicts, sims_dicts, man_dicts, 87.5, 108.0)
            
            # Update peak flags
            peak_map = {r["scan_point_id"]: r["is_peak"] for r in proc if r["scan_point_id"]}
            for p in all_pts:
                if p.id in peak_map: p.is_peak = peak_map[p.id]
                
            id_objs = [
                IdentificationResult(
                    session_id=sess.id,
                    scan_point_id=r["scan_point_id"],
                    frequency_mhz=r["frequency_mhz"],
                    level_dbuvm=r["level_dbuvm"],
                    is_peak=r["is_peak"],
                    identified_name=r["identified_name"],
                    service=r.get("service"),
                    subservice=r.get("subservice"),
                    emission_class=r.get("emission_class"),
                    status=r["status"],
                    distance_km=r.get("distance_km"),
                    station_latitude=r.get("station_latitude"),
                    station_longitude=r.get("station_longitude"),
                    station_city=r.get("station_city"),
                    matched_sims_id=r.get("matched_sims_id"),
                    matched_manual_id=r.get("matched_manual_id"),
                    matched_source=r.get("matched_source", "NONE")
                )
                for r in proc
            ]
            db.bulk_save_objects(id_objs)
            db.commit()
            print("Successfully seeded initial session, scan points, and identification results!")

    except Exception as e:
        print("Seeding note:", e)
    finally:
        db.close()

if __name__ == "__main__":
    seed()
