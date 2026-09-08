import io
from datetime import datetime, date
from typing import List, Dict, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def generate_report_docx(
    session_data: Dict[str, Any],
    results: List[Dict[str, Any]],
    status_summary: Dict[str, int]
) -> bytes:
    """
    Generate professional Balmon Yogyakarta Activity Report in DOCX format.
    """
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Title Header
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = title.add_run("KEMENTERIAN KOMUNIKASI DAN DIGITAL REPUBLIK INDONESIA\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(11)
    r1.font.bold = True
    
    r2 = title.add_run("DIREKTORAT JENDERAL SUMBER DAYA DAN PERANGKAT POS DAN INFORMATIKA\n")
    r2.font.name = "Arial"
    r2.font.size = Pt(10)
    r2.font.bold = True
    
    r3 = title.add_run("BALAI MONITOR SPEKTRUM FREKUENSI RADIO KELAS I YOGYAKARTA\n")
    r3.font.name = "Arial"
    r3.font.size = Pt(11)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(31, 78, 120)
    
    doc.add_paragraph("─" * 60).alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Report Title
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h_run = heading.add_run("LAPORAN HASIL MONITORING & IDENTIFIKASI SPEKTRUM FREKUENSI RADIO")
    h_run.font.name = "Arial"
    h_run.font.size = Pt(13)
    h_run.font.bold = True
    
    # Metadata Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    meta_data = [
        ("Nama Sesi / Kegiatan", session_data.get("session_name", "-")),
        ("Nomor SPT", session_data.get("spt_number", "-")),
        ("Petugas Pelaksana", session_data.get("officer_name", "-")),
        ("Tanggal & Jam Pengamatan", f"{session_data.get('monitoring_date', date.today())} {session_data.get('monitoring_time', '')}"),
        ("Stasiun Monitoring & Lokasi", f"{session_data.get('monitoring_station', 'SMSN')} (Lat: {session_data.get('latitude')}, Lon: {session_data.get('longitude')})"),
        ("Alat Ukur / Format File", f"{session_data.get('device_type', 'TCI')} | Rentang: {session_data.get('start_frequency_mhz', '-')} - {session_data.get('stop_frequency_mhz', '-')} MHz")
    ]
    
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl = row.cells[0]
        cell_val = row.cells[1]
        
        cell_lbl.text = label
        cell_lbl.paragraphs[0].runs[0].font.bold = True
        cell_lbl.paragraphs[0].runs[0].font.size = Pt(9.5)
        cell_lbl.width = Inches(2.2)
        
        cell_val.text = str(val)
        cell_val.paragraphs[0].runs[0].font.size = Pt(9.5)
        cell_val.width = Inches(4.5)
        
    doc.add_paragraph()
    
    # Summary of Status
    doc.add_heading("Rekapitulasi Status Identifikasi", level=2)
    rekap_p = doc.add_paragraph()
    rekap_p.add_run(
        f"Total Teridentifikasi: {len(results)} sinyal | "
        f"Legal: {status_summary.get('Legal', 0)} | "
        f"Kadaluarsa: {status_summary.get('Kadaluarsa', 0)} | "
        f"Belum Diketahui: {status_summary.get('Belum Diketahui', 0)} | "
        f"Ilegal: {status_summary.get('Ilegal', 0)} | "
        f"Off Air: {status_summary.get('Off Air', 0)}"
    )
    
    # Results Table
    doc.add_heading("Tabel Hasil Pengamatan & Identifikasi", level=2)
    
    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    headers = ["No", "Frek (MHz)", "Level (dBuV/m)", "Identifikasi Pengguna", "Status", "Jarak (km)", "Dinas"]
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(8.5)
        # Background color header
        shading = parse_xml(r'<w:shd {} w:fill="1F4E78"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        
    for idx, r in enumerate(results[:200], 1): # Cap at 200 rows for docx readability
        row_cells = table.add_row().cells
        row_cells[0].text = str(idx)
        row_cells[1].text = f"{r.get('frequency_mhz', 0.0):.4f}"
        row_cells[2].text = f"{r.get('level_dbuvm', 0.0):.1f}"
        row_cells[3].text = str(r.get("identified_name") or "-")
        row_cells[4].text = str(r.get("status") or "-")
        row_cells[5].text = f"{r.get('distance_km'):.1f}" if r.get('distance_km') is not None else "-"
        row_cells[6].text = str(r.get("service") or "-")
        
        for c in row_cells:
            c.paragraphs[0].runs[0].font.size = Pt(8)
            
    doc.add_paragraph()
    
    # Signature Section
    doc.add_paragraph()
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    c_left = sig_table.rows[0].cells[0]
    c_right = sig_table.rows[0].cells[1]
    
    p_left = c_left.paragraphs[0]
    p_left.add_run("Mengetahui,\nKoordinator Pengamatan Spektrum\n\n\n\n__________________________")
    
    p_right = c_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.add_run(f"Yogyakarta, {date.today().strftime('%d %B %Y')}\nPetugas Pelaksana Monitoring\n\n\n\n{session_data.get('officer_name', 'Petugas Balmon')}")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
